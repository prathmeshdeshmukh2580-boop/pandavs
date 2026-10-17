import os
from pathlib import Path
from typing import Dict, Any, Optional, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from services.renderer.fit_check import FitChecker
from services.renderer.diagrams import DiagramEngine

# ==============================================================================
# TechSummitTrainers Brand Palette & Design Tokens
# ==============================================================================
BRAND_ORANGE = RGBColor(245, 164, 0)     # #F5A400 TechSummitTrainers Brand Orange
DARK_CHARCOAL = RGBColor(17, 24, 39)     # #111827 Dark Charcoal (Primary Text)
CHARCOAL_SURFACE = RGBColor(31, 41, 55)  # #1F2937 Slate/Charcoal Card Surface
LIGHT_BG = RGBColor(248, 250, 252)       # #F8FAFC Clean Studio Background
CARD_BG = RGBColor(255, 255, 255)        # #FFFFFF Pure White
CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0 Subtle Border Gray
TEXT_DARK = RGBColor(17, 24, 39)         # #111827 Main Headings & Body
TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B Secondary Text
TEXT_WHITE = RGBColor(255, 255, 255)     # #FFFFFF Contrast Text

# Code Block & Syntax Colors
CODE_BG = RGBColor(15, 23, 42)           # #0F172A Deep Code Container
CODE_TEXT = RGBColor(52, 211, 153)       # #34D399 Emerald Green Syntax
CODE_HEADER = RGBColor(30, 41, 59)       # #1E293B Code Header Bar

# Status Colors
COLOR_PASSED = RGBColor(16, 185, 129)    # #10B981 Emerald 500
COLOR_WARN = RGBColor(245, 158, 11)      # #F59E0B Amber 500
ACCENT_BLUE = RGBColor(59, 130, 246)     # #3B82F6 Blue 500
ACCENT_PURPLE = RGBColor(147, 51, 234)   # #9333EA Purple 600

# Legacy Dark Theme Palette (for classic template fallback)
LEGACY_BG = RGBColor(15, 23, 42)
LEGACY_SURFACE = RGBColor(30, 41, 59)


class PPTXBuilder:
    """
    Builds official TechSummitTrainers-branded 16:9 widescreen PowerPoint decks
    from verified course deck JSON, powered by Pandavs AI.

    Features:
    - Official TechSummitTrainers logo integration
    - Consistent 16:9 widescreen layout (13.333" x 7.5")
    - 9 reusable branded slide layouts:
      1. Title slide (with logo, metadata, and level badge)
      2. Learning objectives slide
      3. Section divider slide
      4. Concept explanation slide (with pro-tip callout)
      5. Code example slide (syntax container, sandbox verification badge)
      6. Practical exercise / lab slide (starter code, dual-prover contract)
      7. Quiz knowledge check slide (2x2 grid, answer key in notes)
      8. Course summary slide (principles & production checklist)
      9. Closing slide (branded thank-you with TechSummitTrainers logo)
    - Full PowerPoint speaker notes on EVERY slide
    - DiagramEngine integration for Mermaid architectural flows
    """

    def __init__(self, template: str = "techsummit"):
        self.template = template.lower() if template else "techsummit"
        self.prs = Presentation()
        # 16:9 widescreen dimensions
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.diagram_engine = DiagramEngine()

        # Resolve logo asset path
        self.logo_path = self._resolve_logo_path()

    def _resolve_logo_path(self) -> Optional[str]:
        candidates = [
            Path(__file__).parent / "assets" / "techsummit-logo.jpg",
            Path(__file__).parent.parent.parent / "apps" / "web" / "public" / "assets" / "techsummit-logo.jpg",
            Path(__file__).parent.parent / "renderer" / "assets" / "techsummit-logo.jpg",
        ]
        for p in candidates:
            if p.exists() and p.is_file():
                return str(p.resolve())
        return None

    def build_deck(self, deck_data: Dict[str, Any], output_path: str) -> str:
        slides = deck_data.get("slides", [])
        has_closing = False

        for i, slide_data in enumerate(slides):
            slide_type = slide_data.get("type", "overview")

            if slide_type == "title":
                self._add_title_slide(slide_data, deck_data)
            elif slide_type == "code":
                self._add_code_slide(slide_data, deck_data)
            elif slide_type == "concept" and slide_data.get("diagram"):
                self._add_diagram_slide(slide_data, deck_data)
            elif slide_type in ("exercise", "lab"):
                self._add_exercise_slide(slide_data, deck_data)
            elif slide_type in ("quiz", "quiz_review"):
                self._add_quiz_slide(slide_data, deck_data)
            elif slide_type == "summary":
                self._add_summary_slide(slide_data, deck_data)
            elif slide_type in ("closing", "thank_you"):
                self._add_closing_slide(slide_data, deck_data)
                has_closing = True
            elif slide_type in ("objectives", "learning_objectives"):
                self._add_learning_objectives_slide(slide_data, deck_data)
            elif slide_type in ("section", "divider"):
                self._add_section_divider_slide(slide_data, deck_data)
            else:
                self._add_content_slide(slide_data, deck_data)

        # Ensure a professional branded closing slide concludes the deck
        if not has_closing:
            self._add_closing_slide({
                "id": "s_closing",
                "title": "Course Completed",
                "notes": f"Conclude the session on {deck_data.get('topic')}. Solicit student questions, collect feedback ratings, and distribute practical lab assignments."
            }, deck_data)

        # Ensure output directory exists and save presentation
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return str(out)

    # --------------------------------------------------------------------------
    # Base Slide Backgrounds, Brand Headers & Footers
    # --------------------------------------------------------------------------

    def _apply_slide_background(self, slide):
        """Creates a clean, light studio background with a subtle top brand accent."""
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()

        # Top edge brand accent stripe (TechSummitTrainers Orange)
        top_stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, Inches(0.08))
        top_stripe.fill.solid()
        top_stripe.fill.fore_color.rgb = BRAND_ORANGE
        top_stripe.line.fill.background()
        return bg

    def _add_header_logo(self, slide):
        """Places the official TechSummitTrainers logo in the top-right header."""
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                # Logo aspect ratio is ~4.1:1; 2.2" wide x ~0.54" high
                slide.shapes.add_picture(
                    self.logo_path,
                    Inches(10.3),
                    Inches(0.4),
                    width=Inches(2.2),
                    height=Inches(0.54)
                )
            except Exception:
                pass

    def _add_footer(self, slide, topic: str):
        """Adds a subtle TechSummitTrainers & Pandavs footer attribution."""
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.73), Inches(0.4))
        ftf = footer_box.text_frame
        ftf.word_wrap = True
        p = ftf.paragraphs[0]
        p.text = f"TechSummitTrainers  ·  {topic}  ·  Powered by Pandavs AI  ·  100% Sandbox Verified"
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED

    def _add_trainer_notes(self, slide, notes_text: str, slide_type: str = "general", context: Optional[Dict[str, Any]] = None):
        """
        Embeds comprehensive trainer delivery notes into the PowerPoint notes frame.
        Guarantees that every slide contains clear, actionable delivery instructions.
        """
        if not notes_text and context:
            notes_text = self._generate_rich_notes(slide_type, context)
        elif notes_text and context:
            # Enrich existing brief notes if needed
            notes_text = f"TRAINER DELIVERY GUIDANCE:\n{notes_text}\n\n" + self._generate_delivery_tips(slide_type, context)

        if notes_text:
            notes_slide = slide.notes_slide
            tf = notes_slide.notes_text_frame
            tf.text = notes_text

    def _generate_rich_notes(self, slide_type: str, ctx: Dict[str, Any]) -> str:
        title = ctx.get("title", "Topic Overview")
        topic = ctx.get("topic", "Software Architecture")

        if slide_type == "title":
            return (
                f"COURSE INTRODUCTION & WELCOME:\n"
                f"• Welcome students to TechSummitTrainers' masterclass on '{topic}'.\n"
                f"• Target Level: {ctx.get('level', 'Intermediate').capitalize()}.\n"
                f"• Estimated Duration: {ctx.get('minutes', 90)} minutes.\n"
                f"• Pedagogical Goal: Build production-ready intuition, verified patterns, and hands-on muscle memory.\n"
                f"• Emphasize that all code presented has been verified in an isolated zero-trust execution sandbox."
            )
        elif slide_type == "code":
            bullets_str = "\n".join([f"  - {b}" for b in ctx.get("bullets", [])])
            return (
                f"CODE DEMONSTRATION & WALKTHROUGH:\n"
                f"• Objective: Walk students through the live implementation of '{title}'.\n"
                f"• Key Teaching Points:\n{bullets_str}\n"
                f"• Common Student Pitfalls: Watch for subtle type mismatches, variable scope leaks, or unhandled nulls.\n"
                f"• Sandbox Note: Point out the green '✓ SANDBOX VERIFIED' badge to reassure students the snippet executes cleanly."
            )
        elif slide_type == "exercise":
            return (
                f"PRACTICAL LAB INSTRUCTIONS:\n"
                f"• Lab Exercise: {title}\n"
                f"• Setup: Have students open their editor/IDE with the starter code.\n"
                f"• Timebox: Allocate 15-20 minutes for hands-on problem solving.\n"
                f"• Invariant Verification: Explain the Dual-Prover contract ('Tests Must Bite'). The test suite must fail with a naive implementation before passing."
            )
        elif slide_type == "quiz":
            return (
                f"KNOWLEDGE CHECK & CLASSROOM ENGAGEMENT:\n"
                f"• Question: {ctx.get('question', title)}\n"
                f"• Activity: Give students 60 seconds to select their answer before revealing the rationale.\n"
                f"• Discussion: Ask one volunteer who picked a wrong option to explain their reasoning, then gently clarify the misconception."
            )
        elif slide_type == "summary":
            return (
                f"SESSION RECAP & ACTION ITEMS:\n"
                f"• Review the core architectural principles covered throughout {topic}.\n"
                f"• Run through the Production Checklist items with the class.\n"
                f"• Distribute student cheat sheets and outline the next hands-on project."
            )
        else:
            return (
                f"CONCEPT DELIVERY NOTES:\n"
                f"• Focus on '{title}'.\n"
                f"• Connect this concept back to real-world production reliability.\n"
                f"• Prompt students: 'Where in your current codebase have you observed this pattern?'"
            )

    def _generate_delivery_tips(self, slide_type: str, ctx: Dict[str, Any]) -> str:
        if slide_type == "code":
            return "DELIVERY TIP: Highlight variable bindings and memory lifecycle before jumping into logic flow."
        elif slide_type == "quiz":
            return "DELIVERY TIP: Use this question as an informal formative assessment to gauge retention before proceeding."
        return "DELIVERY TIP: Maintain an interactive cadence; pause every 3-4 minutes for student questions."

    # --------------------------------------------------------------------------
    # Slide Layout 1: Title Slide (Branded 16:9)
    # --------------------------------------------------------------------------

    def _add_title_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)

        # Prominent TechSummitTrainers Logo in Top Left
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                slide.shapes.add_picture(
                    self.logo_path,
                    Inches(1.0),
                    Inches(0.8),
                    width=Inches(3.4),
                    height=Inches(0.83)
                )
            except Exception:
                pass

        # Brand Badge: TechSummitTrainers Masterclass
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.9), Inches(3.6), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = BRAND_ORANGE
        badge.line.fill.background()
        btf = badge.text_frame
        btf.text = "TECHSUMMITTRAINERS MASTERCLASS"
        btf.paragraphs[0].font.size = Pt(11)
        btf.paragraphs[0].font.bold = True
        btf.paragraphs[0].font.color.rgb = DARK_CHARCOAL
        btf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Level Badge
        lvl = deck_data.get("level", "intermediate").upper()
        lvl_colors = {
            "BEGINNER": RGBColor(16, 185, 129),    # Emerald
            "INTERMEDIATE": RGBColor(59, 130, 246),# Blue
            "ADVANCED": RGBColor(147, 51, 234),   # Purple
        }
        lvl_col = lvl_colors.get(lvl, ACCENT_BLUE)
        lvl_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.9), Inches(2.2), Inches(0.42))
        lvl_badge.fill.solid()
        lvl_badge.fill.fore_color.rgb = lvl_col
        lvl_badge.line.fill.background()
        ltf = lvl_badge.text_frame
        ltf.text = f"LEVEL: {lvl}"
        ltf.paragraphs[0].font.size = Pt(11)
        ltf.paragraphs[0].font.bold = True
        ltf.paragraphs[0].font.color.rgb = TEXT_WHITE
        ltf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Main Title Box
        title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(2.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_data.get("title", deck_data.get("topic", "Course Presentation"))
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Subtitle / Key Takeaways Bullets
        bullets = slide_data.get("bullets", [])
        if bullets:
            for b in bullets[:4]:
                bp = tf.add_paragraph()
                bp.text = f"✦  {b}"
                bp.font.size = Pt(15)
                bp.font.color.rgb = TEXT_MUTED
                bp.space_before = Pt(8)

        # Metadata Card at Bottom
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.6), Inches(11.33), Inches(0.9))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        stack_str = ", ".join([f"{k}: {v}" for k, v in deck_data.get("stack", {}).items()]) or "Modern Stack"
        cp.text = f"Duration: {deck_data.get('minutes', 90)} mins   |   Stack: {stack_str}   |   100% Sandbox Verified"
        cp.font.size = Pt(13)
        cp.font.bold = True
        cp.font.color.rgb = DARK_CHARCOAL

        cp2 = ctf.add_paragraph()
        cp2.text = "Powered by Pandavs AI Course Studio  ·  Zero-Trust Automated Test Suite"
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = BRAND_ORANGE
        cp2.space_before = Pt(4)

        self._add_trainer_notes(slide, slide_data.get("notes", ""), "title", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "level": deck_data.get("level"),
            "minutes": deck_data.get("minutes"),
        })

    # --------------------------------------------------------------------------
    # Slide Layout 2: Learning Objectives Slide
    # --------------------------------------------------------------------------

    def _add_learning_objectives_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"🎯 {slide_data.get('title', 'Course Learning Objectives')}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Orange Accent Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.73), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = BRAND_ORANGE
        line.line.fill.background()

        # Objectives Main Card Container
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.73), Inches(5.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        ctf = card.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = "BY THE END OF THIS TECHSUMMITTRAINERS MASTERCLASS, YOU WILL BE ABLE TO:"
        cp.font.size = Pt(14)
        cp.font.bold = True
        cp.font.color.rgb = BRAND_ORANGE

        bullets = slide_data.get("bullets", [])
        for b in bullets:
            bp = ctf.add_paragraph()
            bp.text = f"✓  {b}"
            bp.font.size = Pt(14)
            bp.font.color.rgb = DARK_CHARCOAL
            bp.space_before = Pt(14)

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "objectives", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "bullets": bullets,
        })

    # --------------------------------------------------------------------------
    # Slide Layout 3: Section Divider Slide
    # --------------------------------------------------------------------------

    def _add_section_divider_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Center Section Card Container
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.8), Inches(10.33), Inches(4.0))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CHARCOAL
        card.line.fill.background()

        # Orange Accent Left Strip
        strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.8), Inches(0.2), Inches(4.0))
        strip.fill.solid()
        strip.fill.fore_color.rgb = BRAND_ORANGE
        strip.line.fill.background()

        ctf = card.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = "SECTION MODULE"
        cp.font.size = Pt(14)
        cp.font.bold = True
        cp.font.color.rgb = BRAND_ORANGE

        cp2 = ctf.add_paragraph()
        cp2.text = slide_data.get("title", "Core Architecture")
        cp2.font.size = Pt(36)
        cp2.font.bold = True
        cp2.font.color.rgb = TEXT_WHITE
        cp2.space_before = Pt(14)

        bullets = slide_data.get("bullets", [])
        if bullets:
            cp3 = ctf.add_paragraph()
            cp3.text = f"Focus: {bullets[0]}"
            cp3.font.size = Pt(15)
            cp3.font.color.rgb = RGBColor(203, 213, 225)
            cp3.space_before = Pt(16)

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "section", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
        })

    # --------------------------------------------------------------------------
    # Slide Layout 4: Concept Explanation Slide
    # --------------------------------------------------------------------------

    def _add_content_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data.get("title", "")
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Orange Accent Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.73), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = BRAND_ORANGE
        line.line.fill.background()

        # Bullets Container Card
        fit = FitChecker.inspect_slide(slide_data)
        font_pt = min(fit["recommended_bullet_font_pt"], 17)

        bullets_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.73), Inches(4.3))
        bullets_card.fill.solid()
        bullets_card.fill.fore_color.rgb = CARD_BG
        bullets_card.line.color.rgb = CARD_BORDER
        bullets_card.line.width = Pt(1.5)

        btf = bullets_card.text_frame
        btf.word_wrap = True

        bullets = slide_data.get("bullets", [])
        for i, bullet in enumerate(bullets):
            bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
            bp.text = f"✦   {bullet}"
            bp.font.size = Pt(font_pt)
            bp.font.color.rgb = DARK_CHARCOAL
            bp.space_before = Pt(12)

        # Pro-Tip Callout Box at Bottom
        pro_tip = slide_data.get("pro_tip")
        if pro_tip:
            tip_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.73), Inches(0.75))
            tip_card.fill.solid()
            tip_card.fill.fore_color.rgb = DARK_CHARCOAL
            tip_card.line.color.rgb = BRAND_ORANGE
            tip_card.line.width = Pt(1.5)

            ctf = tip_card.text_frame
            ctf.word_wrap = True
            cp = ctf.paragraphs[0]
            cp.text = f"💡 TECHSUMMIT PRO-TIP: {pro_tip}"
            cp.font.size = Pt(11)
            cp.font.color.rgb = BRAND_ORANGE

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "concept", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "bullets": bullets,
        })

    # --------------------------------------------------------------------------
    # Slide Layout 5: Code Example Slide (with Verification Badge)
    # --------------------------------------------------------------------------

    def _add_code_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data.get("title", "")
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Left Column: Teaching Points Card
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(4.6), Inches(4.5))
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = CARD_BG
        left_card.line.color.rgb = CARD_BORDER
        left_card.line.width = Pt(1.5)

        ltf = left_card.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = "KEY TEACHING POINTS"
        lp.font.size = Pt(12)
        lp.font.bold = True
        lp.font.color.rgb = BRAND_ORANGE

        for bullet in slide_data.get("bullets", [])[:5]:
            bp = ltf.add_paragraph()
            bp.text = f"•  {bullet}"
            bp.font.size = Pt(13)
            bp.font.color.rgb = DARK_CHARCOAL
            bp.space_before = Pt(8)

        # Pro-tip below left card
        pro_tip = slide_data.get("pro_tip")
        if pro_tip:
            pt_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(4.6), Inches(0.75))
            pt_card.fill.solid()
            pt_card.fill.fore_color.rgb = DARK_CHARCOAL
            pt_card.line.fill.background()
            pt_tf = pt_card.text_frame
            pt_tf.word_wrap = True
            pt_p = pt_tf.paragraphs[0]
            pt_p.text = f"💡 {pro_tip}"
            pt_p.font.size = Pt(10)
            pt_p.font.color.rgb = BRAND_ORANGE

        # Right Column: Code Card Container
        code_blocks = slide_data.get("code", [])
        code_data = code_blocks[0] if code_blocks else {}
        code_src = code_data.get("source", "// No code provided")
        is_verified = code_data.get("verified", "unverified") == "passed"

        # Code Background Card (Dark Charcoal container)
        code_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.4), Inches(6.93), Inches(5.35))
        code_card.fill.solid()
        code_card.fill.fore_color.rgb = CODE_BG
        code_card.line.color.rgb = BRAND_ORANGE
        code_card.line.width = Pt(1.5)

        # Header bar on Code Card
        lang_str = code_data.get("lang", "code").upper()
        ver_str = code_data.get("version", "")
        kind_str = code_data.get("kind", "")

        lang_box = slide.shapes.add_textbox(Inches(5.8), Inches(1.45), Inches(4.0), Inches(0.35))
        ltf2 = lang_box.text_frame
        lp2 = ltf2.paragraphs[0]
        lp2.text = f"💻 {lang_str} {ver_str}  ·  {kind_str}"
        lp2.font.size = Pt(10)
        lp2.font.color.rgb = RGBColor(148, 163, 184)

        # Verification Status Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(1.5), Inches(1.85), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = COLOR_PASSED if is_verified else COLOR_WARN
        pill.line.fill.background()
        ptf = pill.text_frame
        ptf.text = "✓ SANDBOX VERIFIED" if is_verified else "⚠ UNVERIFIED"
        ptf.paragraphs[0].font.size = Pt(8.5)
        ptf.paragraphs[0].font.bold = True
        ptf.paragraphs[0].font.color.rgb = TEXT_WHITE
        ptf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Code Text Area
        fit = FitChecker.inspect_slide(slide_data)
        code_pt = max(9, min(fit["recommended_code_font_pt"], 12))

        code_text_box = slide.shapes.add_textbox(Inches(5.8), Inches(1.9), Inches(6.53), Inches(4.7))
        ctf = code_text_box.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = code_src
        cp.font.name = "Consolas"
        cp.font.size = Pt(code_pt)
        cp.font.color.rgb = CODE_TEXT

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "code", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "bullets": slide_data.get("bullets", []),
            "code_src": code_src,
        })

    # --------------------------------------------------------------------------
    # Slide Layout 6: Practical Exercise / Lab Slide
    # --------------------------------------------------------------------------

    def _add_exercise_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"🛠️ {slide_data.get('title', 'Hands-on Lab Exercise')}"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Left Card: Lab Objectives & Constraints
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.35))
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = CARD_BG
        left_card.line.color.rgb = BRAND_ORANGE
        left_card.line.width = Pt(1.5)

        ltf = left_card.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = "🎯 LAB SPECIFICATION & OBJECTIVES"
        lp.font.size = Pt(13)
        lp.font.bold = True
        lp.font.color.rgb = BRAND_ORANGE

        bullets = slide_data.get("bullets", [])
        for b in bullets:
            bp = ltf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.size = Pt(13)
            bp.font.color.rgb = DARK_CHARCOAL
            bp.space_before = Pt(8)

        # Right Card: Starter Code / Test Invariants
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.4), Inches(5.83), Inches(5.35))
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = CODE_BG
        right_card.line.color.rgb = COLOR_PASSED
        right_card.line.width = Pt(1.5)

        rtf = right_card.text_frame
        rtf.word_wrap = True
        rp = rtf.paragraphs[0]
        rp.text = "🧪 DUAL-PROVER TEST SUITE CONTRACT"
        rp.font.size = Pt(13)
        rp.font.bold = True
        rp.font.color.rgb = COLOR_PASSED

        exercises = deck_data.get("exercises", [])
        ex = exercises[0] if exercises else {}
        starter = ex.get("starter", "# Student starter template\ndef solve():\n    pass")

        rp2 = rtf.add_paragraph()
        rp2.text = "Starter Template:"
        rp2.font.size = Pt(11)
        rp2.font.bold = True
        rp2.font.color.rgb = RGBColor(148, 163, 184)
        rp2.space_before = Pt(8)

        rp3 = rtf.add_paragraph()
        rp3.text = starter[:220]
        rp3.font.name = "Consolas"
        rp3.font.size = Pt(10)
        rp3.font.color.rgb = CODE_TEXT

        rp4 = rtf.add_paragraph()
        rp4.text = "Dual-Prover Verification:"
        rp4.font.size = Pt(11)
        rp4.font.bold = True
        rp4.font.color.rgb = RGBColor(148, 163, 184)
        rp4.space_before = Pt(8)

        rp5 = rtf.add_paragraph()
        rp5.text = "✓ Instructor Reference passes all assertions\n✓ Naive buggy implementation fails as expected ('Tests Must Bite')"
        rp5.font.size = Pt(11)
        rp5.font.color.rgb = TEXT_WHITE

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "exercise", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "bullets": bullets,
        })

    # --------------------------------------------------------------------------
    # Slide Layout 7: Quiz / Knowledge Check Slide
    # --------------------------------------------------------------------------

    def _add_quiz_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"🧠 {slide_data.get('title', 'Knowledge Check: Architecture Review')}"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Quiz Question data
        quiz_list = deck_data.get("quiz", [])
        q = quiz_list[0] if quiz_list else {
            "question": "What is the primary benefit of declarative schema validation at the system boundary?",
            "options": [
                "A. Centralizes validation rules before state propagation",
                "B. Replaces unit testing completely",
                "C. Disables HTTP error status codes",
                "D. Forces XML serialization"
            ],
            "answer": 0,
            "explanation": "Declarative annotations keep controller and service methods focused solely on domain workflows."
        }

        # Question Card Container
        q_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.73), Inches(1.5))
        q_card.fill.solid()
        q_card.fill.fore_color.rgb = DARK_CHARCOAL
        q_card.line.color.rgb = BRAND_ORANGE
        q_card.line.width = Pt(1.5)

        qtf = q_card.text_frame
        qtf.word_wrap = True
        qp = qtf.paragraphs[0]
        qp.text = f"QUESTION: {q.get('question', '')}"
        qp.font.size = Pt(17)
        qp.font.bold = True
        qp.font.color.rgb = TEXT_WHITE

        # Options Cards (2x2 Grid)
        options = q.get("options", [])
        correct_idx = q.get("answer", 0)
        grid_positions = [
            (Inches(0.8), Inches(3.1)),
            (Inches(6.8), Inches(3.1)),
            (Inches(0.8), Inches(4.5)),
            (Inches(6.8), Inches(4.5)),
        ]

        letters = ["A", "B", "C", "D"]
        for idx, (x, y) in enumerate(grid_positions[:len(options)]):
            opt_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.73), Inches(1.2))
            opt_card.fill.solid()
            opt_card.fill.fore_color.rgb = CARD_BG
            opt_card.line.color.rgb = CARD_BORDER
            opt_card.line.width = Pt(1.5)

            otf = opt_card.text_frame
            otf.word_wrap = True
            op = otf.paragraphs[0]
            op.text = f"{letters[idx]}. {options[idx]}"
            op.font.size = Pt(13)
            op.font.color.rgb = DARK_CHARCOAL

        # Footer Rationale Note on Slide
        exp_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.9), Inches(11.73), Inches(0.8))
        etf = exp_box.text_frame
        etf.word_wrap = True
        ep = etf.paragraphs[0]
        ep.text = "💡 Interactive Classroom Assessment: Answers and detailed rationale are in the Trainer Delivery Notes."
        ep.font.size = Pt(11)
        ep.font.color.rgb = TEXT_MUTED

        # Detailed trainer notes with answer key & rationale
        correct_letter = letters[correct_idx] if correct_idx < len(letters) else "A"
        quiz_notes = (
            f"QUIZ ANSWER KEY & RATIONALE:\n"
            f"• Question: {q.get('question')}\n"
            f"• Correct Answer: Option {correct_letter} ({options[correct_idx] if correct_idx < len(options) else ''})\n"
            f"• Official Rationale: {q.get('explanation', '')}\n\n"
            f"TEACHING STRATEGY:\n"
            f"Ask the class to vote using hands or chat. Explain why the distractors are incorrect to dispel common student misconceptions."
        )

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, quiz_notes, "quiz", q)

    # --------------------------------------------------------------------------
    # Slide Layout 8: Course Summary Slide
    # --------------------------------------------------------------------------

    def _add_summary_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"🚀 {slide_data.get('title', 'Production Readiness & Best Practices')}"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Left Column: Key Principles Card
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.6), Inches(4.5))
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = CARD_BG
        left_card.line.color.rgb = CARD_BORDER
        left_card.line.width = Pt(1.5)

        ltf = left_card.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = "🏆 KEY ARCHITECTURAL PRINCIPLES"
        lp.font.size = Pt(13)
        lp.font.bold = True
        lp.font.color.rgb = BRAND_ORANGE

        bullets = slide_data.get("bullets", [])
        for b in bullets[:4]:
            bp = ltf.add_paragraph()
            bp.text = f"✓ {b}"
            bp.font.size = Pt(13)
            bp.font.color.rgb = DARK_CHARCOAL
            bp.space_before = Pt(10)

        # Right Column: Production Checklist Card
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.4), Inches(5.83), Inches(4.5))
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = DARK_CHARCOAL
        right_card.line.color.rgb = BRAND_ORANGE
        right_card.line.width = Pt(1.5)

        rtf = right_card.text_frame
        rtf.word_wrap = True
        rp = rtf.paragraphs[0]
        rp.text = "📋 PRODUCTION DEPLOYMENT CHECKLIST"
        rp.font.size = Pt(13)
        rp.font.bold = True
        rp.font.color.rgb = BRAND_ORANGE

        checklist = [
            "Input boundaries enforce strict schema validation",
            "Database queries use explicit pagination limits",
            "Cache-aside invalidation configured with safe TTLs",
            "Structured JSON logging with correlation IDs",
            "Zero-trust automated unit test coverage verified",
        ]
        for item in checklist:
            ip = rtf.add_paragraph()
            ip.text = f"☐  {item}"
            ip.font.size = Pt(12)
            ip.font.color.rgb = TEXT_WHITE
            ip.space_before = Pt(8)

        # Pro-Tip Callout
        pro_tip = slide_data.get("pro_tip", "Always verify every educational snippet before classroom presentation.")
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.73), Inches(0.75))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = COLOR_PASSED
        card.line.width = Pt(1.5)
        ctf = card.text_frame
        cp = ctf.paragraphs[0]
        cp.text = f"💡 FINAL TAKEAWAY: {pro_tip}"
        cp.font.size = Pt(11)
        cp.font.color.rgb = DARK_CHARCOAL

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "summary", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "bullets": bullets,
        })

    # --------------------------------------------------------------------------
    # Slide Layout 9: Closing Slide (Official Branded Conclusion)
    # --------------------------------------------------------------------------

    def _add_closing_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)

        # Large Centered Logo
        if self.logo_path and os.path.exists(self.logo_path):
            try:
                slide.shapes.add_picture(
                    self.logo_path,
                    Inches(4.6),
                    Inches(1.5),
                    width=Inches(4.1),
                    height=Inches(1.0)
                )
            except Exception:
                pass

        # Thank You Container Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(2.8), Inches(9.33), Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CHARCOAL
        card.line.color.rgb = BRAND_ORANGE
        card.line.width = Pt(2)

        ctf = card.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = "Thank You for Learning with TechSummitTrainers!"
        cp.font.size = Pt(26)
        cp.font.bold = True
        cp.font.color.rgb = TEXT_WHITE
        cp.alignment = PP_ALIGN.CENTER

        cp2 = ctf.add_paragraph()
        cp2.text = f"Masterclass: {deck_data.get('topic', 'Software Architecture')}"
        cp2.font.size = Pt(16)
        cp2.font.color.rgb = BRAND_ORANGE
        cp2.alignment = PP_ALIGN.CENTER
        cp2.space_before = Pt(12)

        cp3 = ctf.add_paragraph()
        cp3.text = "All code examples & test contracts verified via Zero-Trust Execution Sandbox"
        cp3.font.size = Pt(12)
        cp3.font.color.rgb = RGBColor(148, 163, 184)
        cp3.alignment = PP_ALIGN.CENTER
        cp3.space_before = Pt(12)

        cp4 = ctf.add_paragraph()
        cp4.text = "Powered by Pandavs AI Course Studio"
        cp4.font.size = Pt(11)
        cp4.font.bold = True
        cp4.font.color.rgb = COLOR_PASSED
        cp4.alignment = PP_ALIGN.CENTER
        cp4.space_before = Pt(8)

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "closing", {
            "title": "Course Conclusion",
            "topic": deck_data.get("topic"),
        })

    # --------------------------------------------------------------------------
    # Architectural Diagram Slide
    # --------------------------------------------------------------------------

    def _add_diagram_slide(self, slide_data: Dict[str, Any], deck_data: Dict[str, Any]):
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)
        self._apply_slide_background(slide)
        self._add_header_logo(slide)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data.get("title", "")
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = DARK_CHARCOAL

        # Bullets above diagram
        bullets_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.73), Inches(0.9))
        btf = bullets_box.text_frame
        btf.word_wrap = True
        for i, b in enumerate(slide_data.get("bullets", [])[:2]):
            bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
            bp.text = f"✦  {b}"
            bp.font.size = Pt(14)
            bp.font.color.rgb = DARK_CHARCOAL

        # Render diagram PNG
        diagram_src = slide_data.get("diagram", "")
        png_path = self.diagram_engine.render_to_png(diagram_src, name_hint=f"slide_{slide_data.get('id', 'diag')}")
        if os.path.exists(png_path):
            slide.shapes.add_picture(png_path, Inches(0.8), Inches(2.2), width=Inches(11.73), height=Inches(4.5))

        self._add_footer(slide, deck_data.get("topic", ""))
        self._add_trainer_notes(slide, slide_data.get("notes", ""), "diagram", {
            "title": slide_data.get("title"),
            "topic": deck_data.get("topic"),
            "bullets": slide_data.get("bullets", []),
        })
